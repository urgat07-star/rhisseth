"""Create the spreadsheet and release reports from sanitized runner evidence."""
import collections
import csv
import datetime as dt
import json
from pathlib import Path
import sys
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'app/backend'))
from hex_rules import coordinates, FIELDS
source=json.loads((ROOT/'reports/hex-metadata-source.json').read_text(encoding='utf-8'))
by_coordinate={(int(row['Q']),int(row['R'])):row for row in source}
rows=[by_coordinate[c] for c in coordinates()]
for row in rows:
    row.setdefault('Статус данных','Параметры из существующей БД' if row.get('Категория') else 'Требует описания')
unknown=[row for row in rows if not row.get('Категория')]
counts=collections.Counter(row.get('Категория','Не определена') for row in rows)
headers=['Q','R','Статус данных',*FIELDS]
date='2026-09-18'
prefix=ROOT/'reports'/('rhisseth-hexes-'+date)
def value(row,field):
    result=row.get(field,'')
    if field in ('Q','R'):return int(result)
    if field=='Состав ландшафта' and result:
        result=' / '.join(f'{part["name"]} {part["percent"]}%' for part in json.loads(result))
    if str(result).startswith(('=','+','-','@','\t','\r')):return "'"+str(result)
    return result
with prefix.with_suffix('.csv').open('w',encoding='utf-8-sig',newline='') as stream:
    writer=csv.writer(stream,delimiter=';');writer.writerow(headers)
    writer.writerows([value(row,field) for field in headers] for row in rows)
book=Workbook();sheet=book.active;sheet.title='Гексы';sheet.append(headers)
for row in rows:sheet.append([value(row,field) for field in headers])
sheet.freeze_panes='D2'
table=Table(displayName='RhissethHexes',ref=f'A1:{sheet.cell(1,len(headers)).column_letter}{len(rows)+1}')
table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);sheet.add_table(table)
for column,field in enumerate(headers,1):sheet.column_dimensions[sheet.cell(1,column).column_letter].width=12 if field in ('Q','R') else 30
for row_index,row in enumerate(rows,2):
    if not row.get('Категория'):
        for cell in sheet[row_index]:cell.fill=PatternFill('solid',fgColor='FFF2CC')
info=book.create_sheet('Сведения');info.append(['Параметр','Значение'])
for entry in [('Дата выгрузки (Europe/Moscow)',date),('Источник','PostgreSQL работающего rhisseth.ru через STU-AUTOMATION-01'),('Гексов видимой сетки',465),('Без установленной категории',len(unknown)),('Изображение','Artistic V6, не изменено'),('Ограничение','Выгружены существующие параметры. Сверка каждого гекса с художественным изображением не выполнена.'),('Ландшафт','Новые доли задаются администраторами; проценты не получены автоматически из PNG.')]:info.append(entry)
info.column_dimensions['A'].width=35;info.column_dimensions['B'].width=110
book.save(prefix.with_suffix('.xlsx'))
check=load_workbook(prefix.with_suffix('.xlsx'));assert check['Гексы'].max_row==466
logs=sorted((ROOT/'reports/automation-logs'/date).glob('*rhisseth-runner-native.md'))
links='\n'.join(f'- [{log.name}](automation-logs/{date}/{log.name})' for log in logs)
technical=f'''# Rhisseth: таблица гексов, владельцы и начальные баронии

Дата: {date}. Часовой пояс: Europe/Moscow; временные метки UTC и Moscow сохранены в журналах.
Область: rhisseth.ru, карта Artistic V6, 3200×2200, радиус гекса 80, действующая сетка без изменений.

## Результат и данные

Все 465 гексов видимой сетки представлены в PostgreSQL и административной таблице. До миграции хранилось 677 записей, после — 723; 258 записей вне видимой сетки сохранены, но исключены из интерфейса и итоговой таблицы.
Категории видимой сетки: Суша — {counts['Суша']}, Побережье — {counts['Побережье']}, Море — {counts['Море']}. У {len(unknown)} гексов категория не установлена: они отмечены «Требует описания». Отсутствие данных не считается четвёртой категорией.
Существующие категории приведены к правилу: 0% суши — море, 100% — суша, промежуточное значение — побережье. Признак «Остров: Да» относит гекс к побережью независимо от доли суши.
Экспорт: [Excel](rhisseth-hexes-{date}.xlsx), [CSV](rhisseth-hexes-{date}.csv). Это снимок параметров PostgreSQL, а не результат автоматического распознавания PNG.

## Поведение и права

- Карточка гекса уменьшена до 285 px, характеристики доступны для просмотра.
- Географическое название меняют admin/moderator отдельной кнопкой и с подтверждением. Группа выбранных гексов может получить общее название.
- Название территории отделено от географического. Игрок меняет название только своей территории. Переименование начальной баронии распространяется на её гексы, которые всё ещё принадлежат игроку.
- Административный интерфейс: https://rhisseth.ru/admin/hexes; доступ admin/moderator проверяется сервером. Поля характеристики изменяются только через явный редактор с подтверждением.
- Типы владельцев: игрок, компьютерное владение, ничейная территория. В ничейной территории жители могут быть описаны в комментарии.
- Компьютерное владение назначается атомарно 2–10 связным соседним гексам. Баронство: 2–3; княжество: 4–6; королевство: 7–10. Название и текст правил сохраняются отдельно.
- Состав ландшафта хранит несколько именованных частей с долями, сумма 100%. Форма позволяет добавлять части без JSON. В карточке можно скрыть проценты; предпочтение сохраняется в браузере.
- Параметр «Пресная вода» удалён из текущих данных, формы, API редактирования и экспортов.
- Старт игры: титул «Барон», одна начальная барония из 2–3 соседних гексов. Можно выбрать предложенный вариант или получить случайный. Выбор ограничен свободной сушей/побережьем с установленной категорией. Повторный старт отклоняется; конкурентная выдача сериализована транзакционной блокировкой.
- Административные формы и карточка передают ревизию данных; устаревшее сохранение возвращает 409. CSRF и серверная проверка ролей сохранены.

## Доказательства и последовательность

09:21–09:29 UTC: просмотр состояния сайта через закреплённый SSH runner, получение текущих параметров и контрольных сумм. Выявлена карта V6 с локальными серверными изменениями; они сохранены.
09:32 UTC: резервирование PostgreSQL и файлов приложения; восстановление дампа в отдельную БД rhisseth_hex_metadata_check; миграция и интеграционные проверки; удаление тестовой БД; применение миграции 003_hex_metadata.sql к рабочей БД; перезапуск только приложения.
09:33 UTC и позднее: проверка работающих API и публичного HTTPS; окончательные метки и команды — в журналах.
Контрольная сумма активного изображения V6 до и после: `b3c27fb8b7ac9d29bd34edcd2148ed65fa62847f5f3ad33cb17da98e344b7fc8`.
Резервная копия первой публикации: /opt/rhisseth/backups/hex-metadata-20260918-093224/database.dump; прежние файлы находятся рядом в files/. Последняя публикация и её backup указаны в соответствующем журнале.

## Проверки

16 профильных unit/API тестов проходят (baseline, full_canvas, site_auth, hex_metadata).
Браузер Edge/Playwright: компактная карточка, просмотр, скрытие долей, отмена/подтверждение названия, форма частей ландшафта, права игрока на название территории, варианты начальной баронии, мобильная ширина; JavaScript ошибок нет.
Интеграционные проверки на восстановленной БД: 465 гексов, удаление поля пресной воды, выбранный старт, отказ повторного старта, переименование баронии, отказ чужому владельцу, независимость двух названий, состав ландшафта, категория острова, компьютерное владение, CSV.
Рабочие API /api/hexes, /api/admin/owners, /api/admin/hexes/export, /api/start/options и административная страница отвечают 200 под проверочной сессией администратора; временная сессия удалена. Публичная HTTPS-проверка выполнена с проверкой сертификата.

## Ограничения, выводы и рекомендации

Параметры из прежней БД не были заново сверены с каждым гексом художественной карты V6. Из изображения нельзя достоверно установить игровые числовые характеристики и ресурсы. Для {len(unknown)} гексов требуется описание; существующие гексы также требуют визуальной проверки при изменении изображения.
Неизвестные значения не заменены вымышленными. Новые доли ландшафта и признаки островов необходимо внести через административную таблицу.
Правила компьютерных владений сохранены как текст; автоматическое поведение компьютерных игроков и симуляция жителей в этот выпуск не входят.
Групповое географическое переименование выполняется последовательно; при ошибке интерфейс сообщает количество уже сохранённых гексов. Назначение компьютерного владения и выдача баронии атомарны.
Код опубликован проверенным пакетом поверх существующего checkout, без Git push и без перезаписи неизвестных изменений. SHA исходной базы checkout: 256e594cde6eda29c3460035ca0d7435c540e51b. Пакет и время отражены в /opt/rhisseth/release.json. Перед следующим Git deployment необходимо включить текущие изменения и карту V6 в согласованный выпуск.
Рекомендуется заполнить гексы без описания, проверить составные ландшафты и острова по изображению, затем определить игровые правила компьютерных владений.
Перезагрузка сервера и PostgreSQL не выполнялась. Изменены схема/данные и код; приложение rhisseth перезапущено. Секреты в отчёты не включены.

## Санитизированные журналы

{links}
'''
(ROOT/'reports'/f'rhisseth-hex-metadata-{date}.md').write_text(technical,encoding='utf-8')
document=Document();document.add_heading('Rhisseth: обновление карты и управления территориями',0)
document.add_paragraph(f'Дата: {date}. Статус: установлено на rhisseth.ru, проверки пройдены.')
document.add_heading('Что изменилось',1)
document.add_paragraph('Создана отдельная административная таблица всех 465 видимых гексов с выгрузкой Excel/CSV. Карточка уменьшена и показывает сведения. Администраторы и модераторы меняют географические названия; владельцы — названия своих территорий. Изображение карты сохранено.')
document.add_paragraph('Установлены три категории: суша, побережье, море. Поддерживаются составные ландшафты с долями и скрытием процентов. Поле пресной воды удалено. Добавлены владельцы: игроки, компьютерные владения, ничейные территории.')
document.add_paragraph('Игрок начинает с титулом барона и получает одну баронию из 2–3 соседних гексов случайно либо выбирает из доступных вариантов. Компьютерные владения получают размерные типы: баронство 2–3, княжество 4–6, королевство 7–10 гексов.')
document.add_heading('Качество данных и требуемые действия',1)
document.add_paragraph(f'У {len(unknown)} гексов отсутствует установленное описание категории. Они выделены в таблице. Прежние параметры получены из действующей базы; полная визуальная сверка карты V6 ещё не выполнена. Нужна работа администраторов по заполнению и проверке местности, долей и островов.')
document.add_paragraph('Текст правил компьютерных владений можно сохранять, но поведение компьютерных игроков ещё необходимо определить и реализовать.')
document.add_heading('Проверки и влияние на сервис',1)
document.add_paragraph('Права доступа и формы проверены автоматическими тестами и в браузере. Перед изменением выполнено резервирование и проверка на восстановленной тестовой базе. Сайт доступен по HTTPS; карта не изменена. Перезапущено только приложение, сервер и база данных не перезагружались.')
document.save(ROOT/'reports'/f'rhisseth-hex-metadata-{date}.docx')
assert len(list(csv.reader(prefix.with_suffix('.csv').open(encoding='utf-8-sig'),delimiter=';')))==466
assert all('Пресная вода' not in r for r in rows)
print(json.dumps({'visible_hexes':len(rows),'without_category':len(unknown),'categories':dict(counts),'excel_rows':check['Гексы'].max_row-1,'reports':'Markdown + Word; Excel + CSV'},ensure_ascii=True))
